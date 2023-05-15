"""
    openpyxl绘制折线图
"""
# 导入模块
import os
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

if __name__ == '__main__':
    # 定义目标文件夹路径
    path_source = '../operatior_data/codes/material/各部门利润表汇总-副本/'
    # 获取存储所有文件名的列表
    list_file_name = os.listdir(path_source)
    # 循环遍历列表
    for file_name in list_file_name:
        # 拼接文件路径
        file_path = path_source + file_name
        # 打开工作簿
        work_book = load_workbook(file_path)
        # 获取活动工作表
        work_sheet = work_book.active
        # 实例化LineChart类
        chart_line = LineChart()
        # 引用工作表数据, 创建Reference对象
        reference = Reference(worksheet=work_sheet, min_row=3, max_row=9, min_col=1, max_col=5)
        # 添加引用数据到LineChart对象中, 创建图表
        chart_line.add_data(reference, from_rows=True, titles_from_data=True)
        # 设置折线图在工作表中的位置
        work_sheet.add_chart(chart_line, "C12")
        # 引用工作表表头数据设置类别轴标签
        chart_line.set_categories(Reference(work_sheet, min_row=2, max_row=2, min_col=2, max_col=5))
        # 设置 x 轴标题
        chart_line.x_axis.title = '季度'
        # 设置 y 轴标题
        chart_line.y_axis.title = '利润'
        # 修改折线图样式
        chart_line.style = 24
        # 设置折线图大小
        chart_line.width = 30
        chart_line.height = 20
        # 保存工作簿
        work_book.save(file_path)
    print('工作表折线图已绘制完毕！！')
